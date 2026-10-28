from pathlib import Path
import csv,json,hashlib,zipfile,datetime,re
from lxml import html
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo

OUT=Path('user/workspace/co-scientist-outputs/upwork-policy-records')
TMP=Path('user/user-data/tmp/upwork-curation')
NOW=datetime.datetime.now(datetime.timezone.utc).isoformat()
sources=[]; evidence={}; rows=[]
def add_source(sid,path,repository,url,title,kind,created='',updated='',capture='',publication='',period=''):
    b=path.read_bytes(); member='sources/'+path.name; evidence[member]=b
    sources.append(dict(source_id=sid,repository=repository,title=title,source_type=kind,source_url=url,original_url=url,created_at=created,updated_at=updated,publication_date=publication,archive_capture_at=capture,reporting_period=period,retrieved_at=NOW,http_status=200,sha256=hashlib.sha256(b).hexdigest(),evidence_member=member,access_note='Substantive primary content retrieved; not a search snippet',rights_note='Publicly accessible; no blanket reuse licence inferred'))
    return sources[-1]
ids=[5313574196627,1500007569061,211062538,360048105134,17975121967891,25205969832083,48889541166483]
texts={}
for i in ids:
    p=TMP/f'{i}.json'; a=json.loads(p.read_text())['article']; sid='UW-'+str(i)
    add_source(sid,p,'Upwork Help Center (Zendesk API)',f'https://support.upwork.com/api/v2/help_center/en-us/articles/{i}.json',a['title'],'current_policy_article',a['created_at'],a['updated_at'])['original_url']=a['html_url']
    texts[sid]=html.fromstring(a['body']).text_content()
for key,cap in [('historical_pre','20230327221055'),('historical_post','20230506084439')]:
    sid='IA-'+cap; url=f'https://web.archive.org/web/{cap}id_/https://support.upwork.com/hc/en-us/articles/211062538-Freelancer-Service-Fees'
    add_source(sid,TMP/f'{key}.html','Internet Archive Wayback Machine',url,'Freelancer Service Fees','archived_policy_article',capture=datetime.datetime.strptime(cap,'%Y%m%d%H%M%S').isoformat()+'Z')['original_url']='https://support.upwork.com/hc/en-us/articles/211062538-Freelancer-Service-Fees'
    texts[sid]=html.fromstring((TMP/f'{key}.html').read_bytes()).text_content()
add_source('SEC-2023',TMP/'sec_2023.html','SEC EDGAR','https://www.sec.gov/Archives/edgar/data/1627475/000162747524000012/upwk-20231231.htm','Upwork Inc. FY2023 Form 10-K','corporate_regulatory_filing',publication='2024-02-15',period='2023')
texts['SEC-2023']=html.fromstring((TMP/'sec_2023.html').read_bytes()).text_content()
add_source('UW-TR2024',TMP/'transparency_2024.pdf','Upwork Help Center attachments','https://support.upwork.com/hc/article_attachments/49240416097939','2024 Upwork Marketplace Transparency Report','corporate_transparency_report',period='2024')
S={x['source_id']:x for x in sources}
def record(sid,topic,name,summary,quote,locator,actor='both',value='',unit='',start='',end='',date_basis='',uncertainty='',rtype='policy_rule',denominator=''):
    s=S[sid]
    flags=[]
    if s['source_type']=='current_policy_article':flags+=['CURRENT_VERSION_NOT_HISTORICAL','EFFECTIVE_DATE_NOT_STATED']
    if s['source_type']=='archived_policy_article':flags+=['CAPTURE_NOT_PUBLICATION_DATE']
    if s['source_type']=='corporate_transparency_report':flags+=['SELF_REPORTED_AGGREGATE','PUBLICATION_DATE_UNKNOWN','NO_CLIENT_FREELANCER_SPLIT','FULL_YEAR_POSTDATES_PDF']
    if uncertainty:flags+=uncertainty.split('|')
    if sid in texts:
        norm=lambda x:re.sub(r'\s+',' ',x.replace('\u00a0',' ')).strip()
        assert norm(quote) in norm(texts[sid]),(name,quote)
    rows.append(dict(record_id=f'UPW-{len(rows)+1:03d}',record_type=rtype,policy_topic=topic,record_name=name,actor_scope=actor,summary=summary,value=value,unit=unit,denominator_scope=denominator,effective_start=start,effective_end=end,date_basis=date_basis or ('Observed current version at retrieval; onset unknown' if sid.startswith('UW-') and sid!='UW-TR2024' else ''),period_start='2024-01-01' if sid=='UW-TR2024' else '',period_end='2024-12-31' if sid=='UW-TR2024' else '',source_id=sid,source_locator=locator,source_quote=quote,source_url=s['source_url'],original_url=s['original_url'],source_updated_at=s['updated_at'],archive_capture_at=s['archive_capture_at'],retrieved_at=NOW,quality_tier='A_PRIMARY_RETRIEVED',uncertainty_flags='|'.join(dict.fromkeys(flags)),provenance_sha256=s['sha256']))
pre='IA-20230327221055';post='IA-20230506084439'
record(pre,'fees','Historical sliding fee tiers','Legacy lifetime-billings-per-client tiers; not the general July 2024 rule.','20% freelancer service fee applied to earnings','Prior to May 3, 2023 / tier bullets','freelancer','20;10;5','percent',end='2023-05-02',date_basis='Article explicitly labels prior-to-May-3 regime',uncertainty='HISTORICAL_ONSET_UNKNOWN|LEGACY_EXCEPTIONS',denominator='Per client lifetime billings: USD 0–500; 500.01–10000; above 10000')
record(pre,'fees','Flat freelancer fee announced','10% general fee for new contracts, replacing tiered fees.','Beginning May 3, 2023, Pacific Time, the freelancer service fees will be 10% for all new contracts regardless of contract value.','Opening note','freelancer',10,'percent',start='2023-05-03',date_basis='Explicit effective day; Pacific Time',uncertainty='SPECIAL_CONTRACT_EXCEPTIONS|END_DATE_NOT_ESTABLISHED')
record(pre,'fees','Legacy 5% grandfathering','Already qualifying contracts retained 5% through 2023.','Contracts that already qualify for a 5% freelancer service fee will continue at 5% through the end of 2023.','Opening note','freelancer',5,'percent',end='2023-12-31',date_basis='Explicit end-of-2023 grandfathering',uncertainty='ELIGIBLE_CONTRACTS_ONLY')
record(pre,'fees','Client contract initiation fee','One-time fee up to USD 4.95 per new contract.','Clients will pay a one-time contract initiation fee of up to $4.95 at the start of each new contract beginning April 26, 2023','FAQ / 2023 pricing changes','client',4.95,'USD maximum',start='2023-04-26',date_basis='Explicit effective date',uncertainty='NOT_ALL_CONTRACT_TYPES|END_DATE_NOT_ESTABLISHED')
record(post,'fees','Flat freelancer fee observed after rollout','Post-change archive confirms general 10% fee; corroboration, not another policy event.','As a freelancer, you’ll pay a 10% freelancer service fee','Opening paragraph','freelancer',10,'percent',start='2023-05-03',date_basis='May 3 change named in legacy-contract section',uncertainty='CORROBORATION_NOT_INDEPENDENT_EVENT|SPECIAL_CONTRACT_EXCEPTIONS')
record('SEC-2023','fees','Fee change confirmed in regulatory filing','FY2023 disclosure confirms tier retirement and 10% replacement.','In May 2023, we retired the tiered service fee structure','Item 7 / Marketplace revenue','freelancer',10,'percent',start='2023-05',date_basis='Month precision only in filing',uncertainty='MONTH_PRECISION|CORROBORATION_NOT_INDEPENDENT_EVENT',rtype='policy_disclosure')
fee='UW-211062538'
record(fee,'fees','Current freelancer fee range','Current general range; exceptions include Enterprise and other contract types.','The fee ranges from 0% to 15% per contract','What is the Freelancer Service Fee?','freelancer','0–15','percent',uncertainty='SPECIAL_CONTRACT_EXCEPTIONS')
record(fee,'fees','Fee locked for contract','Current article states fee is fixed once the contract begins.','Once your contract begins, the fee is fixed and won’t change.','How is the Freelancer Service Fee set?','freelancer')
record(fee,'fees','Fee refund on client refund','Service fee on refunded earnings is returned.','we’ll return the Freelancer Service Fee you paid on that amount.','If money is refunded...','freelancer')
circ='UW-360048105134'
record(circ,'circumvention','Off-platform payment prohibition','Off-platform payment after meeting on Upwork without conversion fee is circumvention.','Circumvention is when a client and freelancer meet on Upwork, but then conduct their relationship outside of Upwork by sending payment through another means without paying the conversion fee.','What is circumvention on Upwork?')
record(circ,'circumvention','Permitted relationship conversion','Conversion-fee-paid relationship is a stated exception.','The Upwork Conversion Fee has been paid for the relationship','Two exceptions / bullet 1')
record(circ,'circumvention','Pre-existing relationship exception','Relationships originating off-platform are a stated exception.','A relationship that started outside the platform, comes onto Upwork, and is then taken outside the platform.','Two exceptions / bullet 2')
record(circ,'circumvention','Stated suspension sanction','Article states suspension and loss of future platform use; not an observed case.','Violating the non-circumvention clause of the Terms of Service will result in account suspension','Important note',rtype='stated_enforcement_procedure')
trust='UW-25205969832083'
record(trust,'identity','Single account / no ban evasion','One account; restricted or blocked users cannot create another.','Use one account. If we restrict or block your account, you can’t create a new one.','The basics')
record(trust,'identity','Account use and subcontracting','Personal account use; subcontracting requires consent and is restricted for licensed work.','Don’t subcontract work without your client\'s consent or if the work requires holding a professional license','Represent yourself authentically','freelancer')
record(trust,'feedback','Genuine feedback rule','Feedback manipulation and misuse prohibited.','You can\'t manipulate or misuse the feedback system in any way.','Be honest')
record(trust,'enforcement','Proportionate response menu','Education, restriction, or block; depends on severity and prior violations.','our response will depend on how serious the issue is and whether you’ve violated the Marketplace Standard before.','Learn about how we enforce these standards',rtype='stated_enforcement_procedure')
record(trust,'enforcement','Suspicion-triggered restriction','Suspected fraud may trigger restriction and further verification.','Your account could also be restricted if we suspect fraud may be taking place.','Learn about how we enforce these standards',rtype='stated_enforcement_procedure')
appeal='UW-5313574196627'
record(appeal,'appeals','Appeal decision criteria','Mistake or new relevant information; disagreement alone unlikely to succeed.','You have new, relevant information that significantly changes the facts originally reviewed','How do I win my appeal?',rtype='stated_enforcement_procedure')
record(appeal,'appeals','Appeal review target','Most appeals reviewed within two business days; not a guaranteed deadline or measured median.','Most appeals are reviewed within two business days.','How long does it take Upwork to review an appeal?',value=2,unit='business days',uncertainty='SERVICE_GUIDANCE_NOT_OBSERVED_METRIC',rtype='stated_enforcement_procedure')
record(appeal,'appeals','Poor-feedback appeal waiting period','Six-month waiting period for suspensions due to ongoing client-feedback concerns.','you’ll need to wait six months before submitting an appeal.','How can I appeal a suspension due to poor feedback?','freelancer',6,'months',rtype='stated_enforcement_procedure')
feedback='UW-17975121967891'
record(feedback,'feedback','Feedback removal grounds','Bribery, pressure, extortion, buying/selling feedback and fake contracts may warrant removal.','This includes bribing, pressuring or extorting another customer to get a specific feedback outcome; buying or selling feedback; and creating fake contracts.','Violations that may warrant feedback removal',rtype='stated_enforcement_procedure')
record(feedback,'feedback','Client feedback response window','Clients must respond within 28 days of contract ending.','You must respond within 28 days of the contract ending','Respond / opening options','client',28,'days after contract end',uncertainty='CLIENT_ARTICLE_SCOPE')
record('UW-1500007569061','reporting','Confidential reporting route','Flag as inappropriate sends report to Trust and Safety; reporter not disclosed.','We will not tell the person you reported that it came from you.','Opening note',rtype='stated_enforcement_procedure')
record('UW-48889541166483','transparency','Transparency report repository','Official repository links 2024 and 2025 transparency reports.','Upwork publishes an annual transparency report detailing how we uphold trust and safety on our platform.','Introduction',rtype='repository_disclosure')
# Values below transcribed from original corporate report prose; pages 3–4 visually checked.
def metric(name,value,unit,page,quote,denom='',uncertainty=''):
    record('UW-TR2024','enforcement' if page<4 else 'appeals',name,name,quote,f'PDF p. {page}',value=value,unit=unit,date_basis='Calendar reporting year 2024; publication date not printed',uncertainty=uncertainty,rtype='reported_aggregate',denominator=denom)
metric('User reports received',501480,'reports',2,'In 2024, we received 501,480 reports of potentially illegal content or violations of our Terms')
metric('User reports acted on',199019,'reports',2,'we took action on 199,019 of them.','501480 user reports; not number of suspended accounts')
metric('Median user flag review time',48,'minutes',2,'The median time it took to review a user flag was 48 minutes.')
metric('Acted-on user-flagged jobs already detected automatically',87,'percent',2,'87% of user-flagged job posts we acted on had already been detected by our automated systems.','User-flagged job posts acted on; denominator count not stated',uncertainty='DENOMINATOR_COUNT_NOT_REPORTED')
metric('DSA trusted-flagger notices received',0,'notices',2,'Upwork did not receive any notices from designated DSA trusted flaggers.')
metric('Job posts removed',827265,'job posts',3,'We removed 827,265 job posts and suspended 184,914 accounts.')
metric('Accounts suspended',184914,'accounts',3,'We removed 827,265 job posts and suspended 184,914 accounts.',uncertainty='UNIQUE_PERSONS_NOT_ESTABLISHED')
metric('Content moderation suspension appeals received',22528,'appeals',4,'In 2024, we received 22,528 total appeals related to content moderation suspensions')
metric('Temporary suspension appeals received',12634,'appeals',4,'12,634 involved temporary suspensions and 9,894 involved permanent suspensions.')
metric('Permanent suspension appeals received',9894,'appeals',4,'12,634 involved temporary suspensions and 9,894 involved permanent suspensions.')
metric('Temporary suspension appeals approved',10790,'appeals',4,'10,790 temporary suspension appeals','12634 temporary suspension appeals',uncertainty='APPROVAL_NOT_EQUIVALENT_TO_ERROR_REVERSAL')
metric('Temporary approvals from required actions completed',96,'percent approximate',4,'approximately 96%','10790 approved temporary suspension appeals',uncertainty='ROUNDED_PERCENTAGE|APPROVAL_NOT_EQUIVALENT_TO_ERROR_REVERSAL')
metric('Temporary approvals from changed enforcement decision',4,'percent approximate',4,'about 4% due to a change in our enforcement decision.','10790 approved temporary suspension appeals',uncertainty='ROUNDED_PERCENTAGE')
metric('Permanent suspension appeals approved',4054,'appeals',4,'4,054 permanent suspension appeals.','9894 permanent suspension appeals',uncertainty='APPROVAL_NOT_EQUIVALENT_TO_ERROR_REVERSAL')
metric('Median suspension appeal review time',6.4,'hours',4,'The median time it took to review these appeals was 6.4 hours.','Content moderation suspension appeals')
record('UW-TR2024','government_requests','Government content removal requests', 'No government requests to remove illegal content reported for 2024.','Upwork did not receive government requests to remove illegal content in 2024.','PDF p. 5',value=0,unit='requests',rtype='reported_aggregate',date_basis='Calendar reporting year 2024')
record('UW-TR2024','dispute_settlement','Notices of out-of-court disputes','No notices of out-of-court disputes reported; not all commercial payment disputes.','Upwork did not receive notice of any out-of-court disputes in 2024.','PDF p. 5',value=0,unit='notices',rtype='reported_aggregate',date_basis='Calendar reporting year 2024',uncertainty='CONTENT_MODERATION_CONTEXT_NOT_PAYMENT_DISPUTES')

log=[
 dict(repository='Upwork Help Center',query='Named policy articles: fees, circumvention, account appeals, feedback, Trust and Safety, DSA',result='7 official article API responses included',limitation='Public-facing HTML endpoints returned 403; official Zendesk JSON endpoints returned 200. API updated_at is not policy effective date.'),
 dict(repository='Internet Archive Wayback Machine',query='CDX exact URL support.upwork.com/hc/en-us/articles/211062538-Freelancer-Service-Fees, years 2023–2024, status 200, month collapse, limit 6',result='2 verified replay snapshots included: 20230327221055 and 20230506084439',limitation='Broad/prefix CDX had 503 errors; selected captures are not an exhaustive version history.'),
 dict(repository='SEC EDGAR',query='CIK 1627475 FY2023 10-K; accession 0001627475-24-000012',result='Original filing HTML retrieved, relevant fee passage checked; filing date checked via SEC submissions metadata by search agent',limitation='Initial unauthenticated request 403; descriptive User-Agent request 200. Disclosure states month, not exact effective day.'),
 dict(repository='Upwork official report attachments',query='2024 Upwork Marketplace Transparency Report linked from DSA article 48889541166483',result='Original 5-page report retrieved; 17 aggregate records transcribed',limitation='Report publication date absent. Corporate self-report, no client/freelancer split or linked case records.'),
 dict(repository='Upwork Legal Center',query='https://www.upwork.com/legal',result='Excluded: HTTP 403',limitation='No binding legal agreement clauses inferred from inaccessible original.'),
 dict(repository='Open web discovery',query='Upwork May 2023 fee change; suspension appeal circumvention feedback; transparency report 2024',result='Used only to locate official repositories; no search snippets used as records',limitation='Secondary fee calculators, blogs, testimonials and anonymous cases excluded.')]

def write_csv(name,data):
    with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
write_csv('upwork_policy_records.csv',rows);write_csv('source_inventory.csv',sources)

definition={
'record_id':'Stable ID in this release; one atomic policy provision or reported aggregate per row.',
'record_type':'policy_rule / policy_disclosure / stated_enforcement_procedure / repository_disclosure / reported_aggregate. Not individual cases.',
'policy_topic':'Controlled topical category; not a judgement about fairness.',
'record_name':'Human-readable provision or metric name.',
'actor_scope':'client, freelancer, or both. Both means policy scope or pooled report, not observed equal treatment.',
'summary':'Curator paraphrase; quote is retained separately.',
'value':'Directly reported value; blank for qualitative rules. No paper-reported values or derived ratios.',
'unit':'Explicit units; percentages use 0–100 scale, not 0–1. USD values nominal.',
'denominator_scope':'Eligibility, tier bands, or denominator; unknown counts explicitly noted.',
'effective_start':'Explicit policy onset, YYYY-MM-DD or YYYY-MM if only month known. Blank means unknown.',
'effective_end':'Explicit historical endpoint where supported; blank does not mean still in force.',
'date_basis':'Why dates were assigned. Pacific Time retained for historical fee change.',
'period_start':'Calendar observation period start for aggregate records; not publication date.',
'period_end':'Calendar observation period end for aggregate records.',
'source_id':'Foreign key to source_inventory.csv and workbook Sources sheet.',
'source_locator':'Article heading, filing section, or PDF page.',
'source_quote':'Short verbatim supporting excerpt, not necessarily the entire surrounding passage.',
'source_url':'Actual retrieved primary endpoint or archive replay.',
'original_url':'Human-facing article or original repository URL.',
'source_updated_at':'Repository-provided metadata; not necessarily substantive revision or effective date.',
'archive_capture_at':'Wayback capture timestamp, not publication date.',
'retrieved_at':'UTC collection time.',
'quality_tier':'A_PRIMARY_RETRIEVED means source content verified; it does not certify corporate claims as unbiased.',
'uncertainty_flags':'Pipe-separated limitations; empty means no additional row-specific flag was necessary.',
'provenance_sha256':'SHA-256 of preserved original source bytes; matches source inventory and evidence ZIP.'}
dictionary=[dict(field=k,definition=v) for k,v in definition.items()]
flagdefs={
'CURRENT_VERSION_NOT_HISTORICAL':'Retrieved current text cannot establish the policy in 2024.',
'EFFECTIVE_DATE_NOT_STATED':'No original onset date assigned from article metadata.',
'CAPTURE_NOT_PUBLICATION_DATE':'Archived timestamp dates capture only.',
'SELF_REPORTED_AGGREGATE':'Corporate-reported enforcement quantity; not independently audited here.',
'PUBLICATION_DATE_UNKNOWN':'2024 report does not print a publication date.',
'NO_CLIENT_FREELANCER_SPLIT':'Report does not split this measure by account role.',
'FULL_YEAR_POSTDATES_PDF':'Full-year 2024 report includes months after supplied July 2024 PDF.',
'HISTORICAL_ONSET_UNKNOWN':'Start of legacy tier schedule not established.',
'LEGACY_EXCEPTIONS':'General legacy bands, not all products or exceptions.',
'SPECIAL_CONTRACT_EXCEPTIONS':'Enterprise and other special contracts may differ.',
'END_DATE_NOT_ESTABLISHED':'No complete subsequent policy history collected.',
'ELIGIBLE_CONTRACTS_ONLY':'Grandfathering not available to all contracts.',
'NOT_ALL_CONTRACT_TYPES':'Special contract exemptions exist; see source.',
'CORROBORATION_NOT_INDEPENDENT_EVENT':'Same change described by another version or disclosure; do not count as separate change.',
'MONTH_PRECISION':'Exact effective day not stated in this filing.',
'SERVICE_GUIDANCE_NOT_OBSERVED_METRIC':'Stated target, not measured service performance.',
'CLIENT_ARTICLE_SCOPE':'Article is directed to clients; not generalized to freelancers.',
'DENOMINATOR_COUNT_NOT_REPORTED':'Percentage denominator described but count unavailable.',
'UNIQUE_PERSONS_NOT_ESTABLISHED':'Accounts are not necessarily distinct people; deduplication method unspecified.',
'APPROVAL_NOT_EQUIVALENT_TO_ERROR_REVERSAL':'Reinstatement may follow remediation, not acknowledgement of erroneous enforcement.',
'ROUNDED_PERCENTAGE':'Report labels percentage as approximate.',
'CONTENT_MODERATION_CONTEXT_NOT_PAYMENT_DISPUTES':'Zero reported out-of-court notices should not be read as zero commercial disputes.'}
assert {f for r in rows for f in r['uncertainty_flags'].split('|') if f} <= flagdefs.keys()
wb=Workbook();wb.remove(wb.active)
for sheet,data in [('Records',rows),('Sources',sources),('Dictionary',dictionary),('Uncertainty flags',[dict(flag=k,meaning=v) for k,v in flagdefs.items()]),('Search log',log)]:
    ws=wb.create_sheet(sheet);ws.append(list(data[0]))
    for d in data:ws.append([d[k] for k in data[0]])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='123C45');c.alignment=Alignment(wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:c.font=Font(name='Arial',size=10);c.alignment=Alignment(vertical='top',wrap_text=True)
        if row[0].row%2==0:
            for c in row:c.fill=PatternFill('solid',fgColor='EDF5F5')
        ws.row_dimensions[row[0].row].height=48
    for col in ws.columns:
        key=col[0].value;ws.column_dimensions[col[0].column_letter].width=20 if key.endswith('_id') else (55 if key in ['summary','source_quote','definition','meaning','limitation'] else 32)
    if sheet=='Records':
        flagcol=list(rows[0]).index('uncertainty_flags')+1
        for cells in ws.iter_rows(min_row=2):
            c=cells[flagcol-1]
            if c.value:c.fill=PatternFill('solid',fgColor='FFF0C2')
wb.save(OUT/'upwork_policy_records.xlsx')
with zipfile.ZipFile(OUT/'source_evidence.zip','w',zipfile.ZIP_DEFLATED) as z:
    for name,b in evidence.items():z.writestr(name,b)
    z.writestr('manifest.json',json.dumps(sources,indent=2))
    z.writestr('search_log.json',json.dumps(log,indent=2))
# Dataset QA: foreign keys, IDs, direct quote checks above, provenance hashes and appeal subtotal consistency.
assert len({r['record_id'] for r in rows})==len(rows)
assert all(r['source_id'] in S for r in rows)
assert 12634+9894==22528
assert 10790<=12634 and 4054<=9894
assert len(load_workbook(OUT/'upwork_policy_records.xlsx')['Records']['A'])==len(rows)+1
with zipfile.ZipFile(OUT/'source_evidence.zip') as z:
    assert all(hashlib.sha256(z.read(s['evidence_member'])).hexdigest()==s['sha256'] for s in sources)
README=f'''# Upwork policy and enforcement records — curated primary-source release

Collected: {NOW}. **{len(rows)} records from {len(sources)} preserved primary-source documents**, across Upwork Help Center/attachments, Internet Archive Wayback Machine and SEC EDGAR. This is a bounded, purposive collection, not an exhaustive repository census. All records are real, directly sourced records; none are simulated.

## What this adds to your paper

Scope follows *Navigating Upwork: Understanding Corporate Dynamics And Unveiling Freelancer Realities Amid Policy Disruption*, DOI 10.2139/ssrn.6031034, especially pp. 2–3, 10, 21–24 and 34–36. The supplied PDF is dated **15 July 2024** throughout, despite the task citation specifying 2026. Its design combines a stated systematic review, user experiences and an autobiographical account. It reports 218 studies/reports and 495 perspectives for one observation and 50 sources for another, but provides no item-level source roster or extractable study dataset. Page 36 says no data beyond public citations were used. Those paper-reported counts and illustrative suspension cases were **not imported as repository records**.

The appropriate record types are dated policy provisions, corporate regulatory disclosures, and corporate-reported enforcement aggregates—not experimental observations or independently adjudicated user cases. The recipe was narrowed accordingly. Trustpilot/forum narratives and generic gig-economy statistics were excluded from the primary dataset; these would require a separate sampling and verification protocol.

## Three consequential observations

1. **Fee chronology needs separation.** The paper presents 20/10/5% tiers as the general schedule (pp. 2, 10, 21). Upwork's March 27, 2023 archived help article explicitly announces a **May 3, 2023 Pacific Time** move to 10% on new contracts, with eligible existing 5% contracts protected through end-2023. A May 6 archive and FY2023 10-K corroborate the change. The current article instead reports a 0–15% range. Its onset was not inferred, and a complete transition history to today's regime is not claimed.
2. **There are usable original enforcement totals.** Upwork's own 2024 transparency report states **184,914 accounts suspended**, **827,265 job posts removed**, and **22,528 suspension appeals received**. It reports 10,790 approved temporary-suspension appeals and 4,054 approved permanent-suspension appeals. These are pooled content-moderation measures, not case-level observations and not a freelancer/client comparison.
3. **Approved appeal ≠ enforcement error.** The report attributes approximately 96% of temporary-suspension appeal approvals to completion of required reinstatement actions, and about 4% to a changed enforcement decision. Treating all approvals as wrongful suspensions would overstate evidence. No causal or fairness conclusion is supported by these aggregates alone.

## Files and structure

- `upwork_policy_records.xlsx`: filterable Records, Sources, Dictionary, Uncertainty flags and Search log sheets; yellow cells highlight limitations.
- `upwork_policy_records.csv`: the same {len(rows)} records, UTF-8 BOM, one atomic provision or reported measure per row.
- `source_inventory.csv`: {len(sources)} source documents with URLs, repository metadata, archive timestamps, access date and SHA-256.
- `source_evidence.zip`: original API JSON, archive HTML, SEC filing HTML and official 2024 report PDF; manifest and search log included.
- `README.md`: scope, findings, quality criteria, date conventions, limitations and validation.

The CSV intentionally mixes numeric measures and qualitative policy provisions using `record_type`; **filter before analysis**. Multiple rows can share a source; corroborating versions must not be counted as separate policy events. Dates are ISO 8601; a blank means unknown/not applicable, never zero. Percentages use the 0–100 scale. Monetary amounts are nominal USD. Exact source quotations and page/section locators remain alongside harmonized paraphrases. No ratios were calculated in the dataset, and no missing values were invented.

## Source selection and access

Included only substantive content retrieved directly from official repositories, or Internet Archive captures of official Upwork pages. Metadata-only or search-snippet matches were insufficient. Upwork's public HTML help/legal pages initially returned HTTP 403; official Zendesk article JSON endpoints successfully exposed the original article bodies and metadata. The Legal Center remained inaccessible and was excluded: support summaries are not represented as a complete binding legal agreement. SEC EDGAR returned 200 with a descriptive User-Agent after an initial 403. Wayback exact-URL lookup succeeded, while broader CDX queries encountered 503 errors. The archive sample comprises two checked snapshots, not all revisions. All eleven source documents are preserved with byte-level hashes.

Quality tier `A_PRIMARY_RETRIEVED` means the original primary content was retrieved and inspected; it is **not** an independent endorsement of corporate reporting accuracy. Search discovery used fee-change, appeals/circumvention/feedback, and transparency-report queries; exact repository queries and exclusions are in Search log and the ZIP. The SEC filing date was checked against accession metadata by the repository-search agent; the filing HTML is preserved. The metadata response itself is not preserved.

## Temporal and interpretive limits

- Historical snapshots, present-day articles and full-year 2024 aggregates remain distinguishable. Current text is **not backdated to July 2024**. Article `created_at` and `updated_at` are repository metadata, not original effective dates.
- The 2024 transparency report does not print its publication date. It was discovered through the current official DSA article; its full-year figures necessarily include months after the supplied PDF's July 2024 date. This is an extension to the paper's evidence base, not evidence the author had available then.
- Report metrics may overlap across reporting, automation and enforcement. Do not add reports, removed jobs and suspended accounts together. Counts of accounts are not proven unique people; no deduplicated case linkage or enforcement-rate denominator is provided.
- The report has rounded graphic breakdowns (e.g. 117k permanent / 68k temporary suspensions). Those were **not substituted for exact prose totals** and are not included as exact counts.
- No role, country, job type or protected-characteristic breakdown supports testing equal treatment. No independent case findings, appeal case files, individual notifications or adjudicated account-level outcomes were collected. Payment-dispute outcomes are not covered by this release; the reported zero out-of-court notices is in a content-moderation context.
- Public access does not imply unrestricted redistribution rights; source copyrights and terms remain applicable. Preserved evidence is for provenance/research review, not an assertion of an open-data licence.

## Validation completed

Unique record IDs, valid source foreign keys, non-PDF quotation substring checks against retrieved originals (whitespace normalized), full source-file SHA-256 verification within the ZIP, CSV/workbook record-count agreement, defined uncertainty flags, and appeal subtotal checks (12,634 + 9,894 = 22,528). The five-page transparency report was read in full; enforcement and appeals pages 3–4 were also checked visually. Exact numeric values were taken from report prose, not rounded chart labels. This is a single-curator release, not duplicate independent screening.

## Reuse and expansion

For historical policy analysis, filter archived/disclosure records and inspect their explicit date precision. For enforcement summaries, filter `reported_aggregate`, retain units and denominator scopes, and do not interpret approvals as errors. Recommended next work: build a 2023–2026 policy version timeline; add the official 2025 report using the same schema; or collect separately coded, verifiable suspension/appeal cases to investigate role-specific treatment.
'''
(OUT/'README.md').write_text(README)
print(json.dumps(dict(records=len(rows),sources=len(sources),aggregate_records=sum(r['record_type']=='reported_aggregate' for r in rows),qa='passed',files=[p.name for p in OUT.iterdir()]),indent=2))
