from pathlib import Path
import csv,json,zipfile,hashlib,datetime,re,textwrap
from lxml import html
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path('user/workspace/co-scientist-outputs'); OLD=ROOT/'upwork-policy-records'; OUT=ROOT/'upwork-policy-timeline'; TMP=Path('user/user-data/tmp/upwork-timeline')
NOW=datetime.datetime.now(datetime.timezone.utc).isoformat(); TODAY=NOW[:10]
with (OLD/'upwork_policy_records.csv').open(encoding='utf-8-sig') as f: records=list(csv.DictReader(f))
with (OLD/'source_inventory.csv').open(encoding='utf-8-sig') as f: sources=list(csv.DictReader(f))
with zipfile.ZipFile(OLD/'source_evidence.zip') as z: evidence={s['evidence_member']:z.read(s['evidence_member']) for s in sources}
for s in sources:s['evidence_representation']='original retrieved source bytes';s['used_for']='Prior policy collection; see Timeline and Undated rules source links'
new=[('IA-20230926143649','sunset_2023.html','https://web.archive.org/web/20230926143649id_/https://support.upwork.com/hc/en-us/articles/211062538','2023-09-26T14:36:49Z','original archive replay HTML'),('IA-20250427142321','transition_2025_excerpt.txt','https://web.archive.org/web/20250427142321id_/https://support.upwork.com/hc/en-us/articles/211062538-Freelancer-Service-Fee','2025-04-27T14:23:21Z','verbatim article-body excerpt returned by web_fetch; not original HTML'),('UW-FEE-'+TODAY,'fee_current.json','https://support.upwork.com/api/v2/help_center/en-us/articles/211062538.json','','original retrieved JSON bytes')]
for sid,fn,url,cap,representation in new:
    b=(TMP/fn).read_bytes();member='timeline_new_sources/'+fn;evidence[member]=b
    s=dict.fromkeys(sources[0],'');s.update(source_id=sid,repository='Upwork Help Center via Internet Archive' if cap else 'Upwork Help Center (Zendesk API)',title='Freelancer Service Fee',source_type='archived_policy_article' if cap else 'current_policy_article',source_url=url,original_url='https://support.upwork.com/hc/en-us/articles/211062538',archive_capture_at=cap,retrieved_at=NOW,http_status=200,sha256=hashlib.sha256(b).hexdigest(),evidence_member=member,evidence_representation=representation,used_for='New timeline evidence',rights_note='Public access; no blanket reuse licence inferred',access_note='web_fetch retrieved substantive content; raw requests returned 429' if fn.endswith('txt') else 'Substantive content retrieved directly')
    if fn.endswith('json'):
        a=json.loads(b)['article'];s.update(created_at=a['created_at'],updated_at=a['updated_at'],original_url=a['html_url'])
    sources.append(s)
S={s['source_id']:s for s in sources}
texts={}
for s in sources:
    b=evidence[s['evidence_member']]
    if s['evidence_member'].endswith('json'):texts[s['source_id']]=html.fromstring(json.loads(b)['article']['body']).text_content()
    elif s['evidence_member'].endswith('html'):texts[s['source_id']]=html.fromstring(b).text_content()
    elif s['evidence_member'].endswith('txt'):texts[s['source_id']]=b.decode()
rows=[]
def event(date,label,kind,title,before,after,scope,sid,quote,locator,flags,change_id='',refs=''):
    norm=lambda t:re.sub(r'\s+',' ',t.replace('\u00a0',' ')).strip()
    assert norm(quote) in norm(texts[sid]),(sid,quote)
    s=S[sid];rows.append(dict(timeline_id=f'TL-{len(rows)+1:02d}',date=date,date_label=label,date_type=kind,event=title,before=before,after=after,affected_scope=scope,change_id=change_id,source_id=sid,source_locator=locator,source_quote=quote,source_url=s['source_url'],supporting_source_ids=refs,archive_capture_at=s['archive_capture_at'],retrieved_at=s['retrieved_at'],evidence_representation=s['evidence_representation'],uncertainty_flags=flags,sha256=s['sha256']))
pre='IA-20230327221055';post='IA-20230506084439';sunset='IA-20230926143649';trans='IA-20250427142321';live='UW-FEE-'+TODAY
q='Starting May 1, 2025, the Freelancer Service Fee for new contracts will change from a fixed 10% to a variable fee based on factors such as client demand for certain work.'
event('','Before 3 May 2023','historical_baseline','Sliding-scale freelancer fees','Onset not established','20% / 10% / 5%, by lifetime billings per client','General marketplace: first USD 500; USD 500.01–10000; above USD 10000',pre,'Prior to May 3, 2023, Pacific Time','Historical fee heading','ONSET_UNKNOWN|SPECIAL_CONTRACT_EXCEPTIONS')
event('2023-03-27','27 Mar 2023','announcement_observed_by','2023 fee-change notice preserved','Prior announcement date unknown','Notice announces flat fee and temporary grandfathering','Freelancers/agencies; separately stated client initiation fee',pre,'Beginning May 3, 2023, Pacific Time, the freelancer service fees will be 10% for all new contracts regardless of contract value.','Opening note','ARCHIVE_CAPTURE_NOT_PUBLICATION|NOT_A_CHANGE_DATE')
event('2023-04-26','26 Apr 2023','explicit_effective_date','Client contract initiation fee starts','Prior initiation-fee rule not established here','One-time fee up to USD 4.95 per new contract','Eligible new client contracts; not all product types',pre,'Clients will pay a one-time contract initiation fee of up to $4.95 at the start of each new contract beginning April 26, 2023','FAQ / 2023 pricing changes','SPECIAL_CONTRACT_EXCEPTIONS|LATER_CLIENT_FEE_CHANGES_NOT_COVERED','CH-01')
event('2023-05-03','3 May 2023 (Pacific Time)','explicit_effective_date','General freelancer fee becomes 10%','20% / 10% / 5% sliding tiers','10% new contracts; eligible existing 5% protected to year-end','Standard marketplace; agencies included; product exceptions apply',pre,'Beginning May 3, 2023, Pacific Time, the freelancer service fees will be 10% for all new contracts regardless of contract value.','Opening note and transition FAQ','SPECIAL_CONTRACT_EXCEPTIONS|20_PERCENT_TRANSITION_NEXT_INVOICE','CH-02',post+';SEC-2023')
event('2023-05-06','6 May 2023','implementation_observation','Post-change archive confirms 10%','Announced 10% policy','Archived text now describes 10% general service fee','General marketplace contracts',post,'As a freelancer, you’ll pay a 10% freelancer service fee','Opening paragraph','CORROBORATION_NOT_NEW_CHANGE|ARCHIVE_CAPTURE_NOT_PUBLICATION',refs=pre)
event('2023-12-31','Through 31 Dec 2023','explicit_sunset','Eligible legacy 5% protection ends','Eligible pre-existing contracts retained 5%','Temporary protection lasts through end-2023','Qualifying existing contracts, not all future work with the same client',sunset,'you will continue to pay freelancer service fees at 5% of earnings on that contract through the end of 2023.','Exceptions / Legacy contracts','END_OF_YEAR_LANGUAGE|SPECIAL_CONTRACT_EXCEPTIONS','CH-03',pre+';SEC-2023')
event('2024-01-01','1 Jan 2024 (inferred boundary)','derived_calendar_boundary','First day after legacy protection expires','End-of-2023 grandfathering','General 10% regime applies absent another exception; invoice timing not verified','Previously protected legacy contracts only',sunset,'you will continue to pay freelancer service fees at 5% of earnings on that contract through the end of 2023.','Exceptions / Legacy contracts','CALENDAR_INFERENCE_NOT_LITERAL_DATE|NOT_SEPARATE_CHANGE|INVOICE_TIMING_NOT_VERIFIED',refs=pre)
event('2025-04-27','27 Apr 2025','announcement_observed_by','Variable-fee transition notice preserved','General flat 10%','Notice specifies May 1 start; existing contract fees unchanged','New contracts affected; existing contracts grandfathered',trans,q,'Opening note','ARCHIVE_CAPTURE_NOT_PUBLICATION|NOT_A_CHANGE_DATE|PRESERVED_EXCERPT_NOT_RAW_HTML')
event('2025-05-01','1 May 2025','explicit_effective_date','Variable freelancer fee starts for new contracts','Fixed 10% general schedule','Variable fee; fee stays fixed for life of contract; existing fees unchanged','New contracts, not every active freelancer/client relationship',trans,q,'Opening note','LAUNCH_RANGE_NOT_STATED|PENDING_PROPOSAL_CUTOFF_UNRESOLVED|SPECIAL_CONTRACT_EXCEPTIONS|PRESERVED_EXCERPT_NOT_RAW_HTML','CH-04')
event(TODAY,TODAY,'current_policy_observation','Current article states 0–15% per contract','Variable launch established; full subsequent history not reconstructed','0–15%; current article also describes proposal/offer/contract lock-in','Standard fees with Enterprise and other product exceptions',live,'The fee ranges from 0% to 15% per contract','What is the Freelancer Service Fee?','OBSERVATION_NOT_NEW_CHANGE|RANGE_ONSET_NOT_VERIFIED|CURRENT_WORDING_NOT_BACKDATED',refs='UW-211062538')
undated=[r for r in records if r['record_type'] in ['policy_rule','stated_enforcement_procedure'] and not r['effective_start'] and not r['effective_end'] and r['policy_topic']!='fees']
for r in undated:r.update(timeline_status='Not assigned a historical change date; observed in original collection on 2026-10-23')
fields=[dict(field=k,meaning=v) for k,v in {
'date':'ISO date for this particular event type; blank for baseline of unknown onset.',
'date_type':'explicit_effective_date or explicit_sunset counts as change; observations and derived boundary do not.',
'change_id':'Four distinct documented changes CH-01 through CH-04. Empty for baseline, corroboration and other non-change milestones.',
'before / after':'Qualitative rule comparison; stated only within the documented product/contract scope.',
'affected_scope':'Eligibility and exceptions; grandfathering attaches to contracts, not people or client relationships.',
'source_id / supporting_source_ids':'Foreign keys to Sources. Supporting documents corroborate, not additional independent policy events.',
'source_quote / source_locator':'Verbatim supporting passage and heading; full contextual rule may span more than the excerpt.',
'archive_capture_at':'Preservation timestamp, not policy onset or original publication.',
'evidence_representation':'Original bytes versus preserved verbatim tool-returned excerpt; hashes identify the representation actually saved.',
'uncertainty_flags':'Pipe-delimited limits. No implied policy change from metadata revision alone.',
'Undated rules sheet':f'{len(undated)} provisions observed in the prior collection, not a reconstructed history of enforcement changes.'}.items()]
log=[dict(repository='Original collection',query='Read 42 records and 11 preserved source documents',result='Reuse provenance without altering original files'),dict(repository='Internet Archive',query='Exact historical Upwork article captures 20230926143649 and 20250427142321',result='2023 snapshot preserved as HTML. 2025 substantive notice verified by agent and web_fetch; raw replay requests returned HTTP 429, so preserved excerpt clearly labelled.'),dict(repository='Upwork Zendesk API',query='GET /api/v2/help_center/en-us/articles/211062538.json',result='Retrieved current original JSON; metadata update is not treated as policy change. No proof here of exact 0–15% launch range.'),dict(repository='Scope limits',query='No exhaustive crawl or automated diff of all enforcement policies',result='Only dated changes supported by checked sources included. 2024 enforcement totals are not policy-change events.')]
def csvout(path,data):
    with path.open('w',newline='',encoding='utf-8-sig') as f:w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
csvout(OUT/'policy_change_timeline.csv',rows)
wb=Workbook();wb.remove(wb.active)
for name,data in [('Timeline',rows),('Undated rules',undated),('Sources',sources),('Field guide',fields),('Search log',log)]:
    ws=wb.create_sheet(name);ws.append(list(data[0]))
    for d in data:ws.append([d[k] for k in data[0]])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for row in ws:
        for c in row:c.font=Font(name='Arial',size=10);c.alignment=Alignment(wrap_text=True,vertical='top')
    for c in ws[1]:c.font=Font(name='Arial',bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='143F49')
    for col in ws.columns:
        key=col[0].value;ws.column_dimensions[col[0].column_letter].width=54 if key in ['before','after','source_quote','affected_scope','uncertainty_flags','meaning','result'] else 28
    for i in range(2,ws.max_row+1):
        ws.row_dimensions[i].height=64
        for c in ws[i]:c.fill=PatternFill('solid',fgColor='F1F6F7' if i%2==0 else 'FFFFFF')
    if name=='Timeline':
        for i,r in enumerate(rows,2):
            if r['change_id']:
                for c in ws[i]:c.fill=PatternFill('solid',fgColor='DDF1E6')
wb.save(OUT/'policy_change_timeline.xlsx')
with zipfile.ZipFile(OUT/'timeline_evidence.zip','w',zipfile.ZIP_DEFLATED) as z:
    for name,b in evidence.items():z.writestr(name,b)
    z.writestr('source_manifest.json',json.dumps(sources,indent=2));z.writestr('search_log.json',json.dumps(log,indent=2))

# Chronological schematic: row spacing is intentional, not elapsed-time proportional.
fig,ax=plt.subplots(figsize=(15,11),facecolor='#F7FAFA');ax.set_facecolor('#F7FAFA');ax.set_xlim(0,1);ax.set_ylim(-1,11.8);ax.axis('off')
ax.text(.035,11.35,'UPWORK | POLICY-CHANGE TIMELINE',fontsize=22,fontweight='bold',color='#143F49')
ax.text(.035,10.85,'2023–2026 · Four dated changes, with observations and inferences kept separate',fontsize=12,color='#48616A')
ax.plot([.245,.245],[.25,10.0],color='#B9CDCE',lw=2,zorder=0)
short=[('Historical baseline','20% / 10% / 5% by lifetime client billings. Original onset unknown.'),('Notice preserved','Flat-fee change announced by this capture; publication day unknown.'),('Client initiation fee','Up to USD 4.95 per eligible new contract.'),('10% general freelancer fee','New contracts at 10%; eligible existing 5% protected through 2023.'),('Post-change confirmation','Archive describes the implemented 10% general fee; not another change.'),('Legacy 5% protection expires','End-of-2023 sunset for eligible contracts; special products may differ.'),('First day after sunset','Calendar implication, not a literal effective day verified in source.'),('Variable-fee notice preserved','New contracts affected; existing contract fees remain unchanged.'),('Variable fee starts','New-contract fee varies, then stays fixed for that contract. Launch range unstated.'),('Current 0–15% range observed','Current range confirmed; do not backdate this numerical range to launch.')]
for i,(r,(title,detail)) in enumerate(zip(rows,short)):
    y=9.85-i*1.01;change=bool(r['change_id']);derived=r['date_type']=='derived_calendar_boundary';color='#16805E' if change else '#BD8523' if derived else '#64818B'
    ax.scatter([.245],[y],s=115 if change else 58,color=color,zorder=3,edgecolor='white',linewidth=1.5)
    ax.text(.025,y+.06,r['date_label'].replace(' (Pacific Time)','\nPacific Time').replace(' (inferred boundary)','\ninferred boundary'),fontsize=10.2,color='#304B55',va='center')
    ax.text(.28,y+.17,title,fontsize=12,fontweight='bold',color=color)
    ax.text(.28,y-.13,detail,fontsize=10.1,color='#304B55')
    ax.text(.28,y-.40,r['timeline_id']+' · '+r['source_id'],fontsize=8.3,color='#71838A')
ax.text(.035,-.55,'GREEN = dated policy change   ·   GREY = observation / baseline   ·   AMBER = calendar inference',fontsize=10,color='#304B55')
ax.text(.035,-.89,'Chronological schematic; spacing is not time-proportional. Enforcement rule onsets remain undated. See workbook for quotes and exceptions.',fontsize=9,color='#60767D')
fig.savefig(OUT/'policy_change_timeline.png',dpi=170,bbox_inches='tight',facecolor=fig.get_facecolor());plt.close(fig)
md=f'''# Upwork policy-change timeline

Compiled {TODAY}. **Four distinct dated changes**, within a ten-entry chronology. This is a bounded reconstruction of general marketplace fee changes, with an explicit gap register for enforcement policies. It is not a complete history of Upwork's terms or products.

## Timeline

| Date | Milestone | Evidence status |
|---|---|---|
| Before 3 May 2023 | General freelancer fees follow 20/10/5% lifetime-client-billing tiers. | Historical baseline; original onset unknown. |
| By 27 Mar 2023 | Archived help article announces the 2023 changes. | Capture date, not publication or effective date. |
| **26 Apr 2023** | **Client contract initiation fee up to USD 4.95 starts on eligible new contracts.** | Explicit effective date; subsequent client-fee changes not reconstructed. |
| **3 May 2023, Pacific Time** | **General freelancer fee becomes 10%; eligible existing 5% contracts remain protected through year-end.** | Explicit effective date; 20% contracts move on the next relevant invoice. |
| 6 May 2023 | Archived help article describes the 10% rule after rollout. | Implementation corroboration, not an additional policy change. |
| **Through 31 Dec 2023** | **Temporary legacy 5% protection expires.** | Explicit end-of-year sunset; applies to eligible contracts, not all special-product fees. |
| 1 Jan 2024 | First day after that sunset. | Calendar inference; not a separately verified literal date or invoice date. |
| By 27 Apr 2025 | Archived article announces a variable fee for new contracts, with existing contract fees unchanged. | Notice observed by capture date; publication date unknown. |
| **1 May 2025** | **New contracts move from fixed 10% to variable freelancer fees.** | Explicit effective date; fee remains fixed over each contract's life. |
| {TODAY} | Live article reports a 0–15% per-contract range. | Current observation; exact onset of this numerical range not verified here. |

## The missing transition is now dated

The April 27, 2025 archived Upwork notice states:

> “Starting May 1, 2025, the Freelancer Service Fee for new contracts will change from a fixed 10% to a variable fee based on factors such as client demand for certain work. The fee will stay the same for the life of the contract. Fees on existing contracts will not change.”

This establishes **May 1, 2025**, the new-contract scope and grandfathering of existing contract fees. It does **not** state the numerical launch range. The currently retrieved article supplies 0–15%; this reconstruction deliberately does not combine the two documents into an unsupported claim that 0–15% was necessarily the precise launch range. Current proposal/offer lock-in wording is also not retroactively treated as the verified cutoff rule for pending proposals in May 2025.

## Contract scope matters

The 2023 protection attached to eligible existing contracts already at 5%, not all later work with the same client. The September 26, 2023 archive explicitly says new contracts with that client have a 10% fee regardless of client history. Similarly, the 2025 notice preserves existing contract fees, not a permanent fee entitlement for the freelancer or client relationship. Enterprise, Payroll, Any Hire and other special arrangements may differ; do not generalize the standard fee schedule to them.

## Enforcement-policy gaps

The workbook retains **{len(undated)} undated policy/procedure records** from the original October 23 collection: circumvention, account identity, feedback, enforcement, reporting and appeals. Those sources document the rules at retrieval, not when the rules started or changed. No change date is inferred from article creation or update metadata. Their underlying versions were not all refreshed in this follow-up.

The 2024 transparency report describes outcomes over a year. Its suspension and appeal totals are **not policy-change events**, so they do not appear as timeline changes. A later report or updated help article alone does not establish that the enforcement rule changed. Testing differential treatment requires a separate case-level or disaggregated dataset.

The supplied paper's PDF carries July 15, 2024 dates despite the task citation saying 2026. This timeline provides historical context around that date, but does not resolve bibliographic version history or assume that the later 2025 policy was in force when the PDF was written.

## Sources and evidence handling

The original 42-record dataset is unchanged. This follow-up reuses its eleven preserved source documents and adds three representations: the September 2023 archive HTML, an April 2025 transition-notice excerpt, and refreshed fee-article JSON. All fourteen are described in the workbook Sources sheet and bundled manifest. Not every retained source supports a dated event: some support the undated gap register and prior context.

The April 2025 notice was independently retrieved by the repository-search agent and read through web_fetch. Raw HTML preservation attempts returned HTTP 429. The bundle therefore contains a **labelled verbatim article-body excerpt**, not falsely labelled original HTML. Its SHA-256 identifies that saved excerpt representation. Other source hashes identify preserved original bytes. No rate-limit bypass was attempted. Capture timestamps must not be read as original publication dates.

## Deliverables

- `policy_change_timeline.png`: chronological schematic; green = dated change, grey = observation/baseline, amber = inferred calendar boundary. Row spacing is not elapsed-time proportional.
- `policy_change_timeline.xlsx`: Timeline, Undated rules, Sources, Field guide and Search log sheets. Green rows mark the four distinct dated changes.
- `policy_change_timeline.csv`: ten timeline entries with before/after rules, scope, source quotations, links, date types and limitations.
- `timeline_evidence.zip`: fourteen preserved source representations plus manifest and search log; includes original collection evidence.
- `timeline_notes.md`: this interpretive note and evidence limitations.

## Validation

Unique timeline IDs; exactly four distinct nonempty change IDs; all primary quotations checked against preserved source content with whitespace normalization; source foreign keys validated; all fourteen ZIP hashes checked; workbook and CSV row counts matched. No formulas or inferred numeric rates are used. The capture/announcement/implementation distinctions are encoded explicitly to prevent double-counting policy changes. This is a single-curator reconstruction, not independently duplicated screening or an exhaustive archive crawl.
'''
(OUT/'timeline_notes.md').write_text(md)
assert len(rows)==10 and len({r['change_id'] for r in rows if r['change_id']})==4
assert len({r['timeline_id'] for r in rows})==len(rows)
assert all(r['source_id'] in S for r in rows)
with zipfile.ZipFile(OUT/'timeline_evidence.zip') as z:assert all(hashlib.sha256(z.read(s['evidence_member'])).hexdigest()==s['sha256'] for s in sources)
assert load_workbook(OUT/'policy_change_timeline.xlsx')['Timeline'].max_row==11
with (OUT/'policy_change_timeline.csv').open(encoding='utf-8-sig') as f: assert len(list(csv.DictReader(f)))==len(rows)
print(json.dumps({'timeline_entries':len(rows),'distinct_changes':4,'undated_rules':len(undated),'source_representations':len(sources),'qa':'passed'},indent=2))
